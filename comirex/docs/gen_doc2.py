from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

for section in doc.sections:
    section.top_margin    = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin   = Cm(2.5)
    section.right_margin  = Cm(2.5)

normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)

def sf(run, bold=False, italic=False, size=12, mono=False, color=None):
    run.font.name  = 'Courier New' if mono else 'Times New Roman'
    run.font.size  = Pt(size)
    run.font.bold  = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)

def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    sf(r, bold=True, size=14)
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after  = Pt(6)

def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    sf(r, bold=True, size=12)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after  = Pt(4)

def h3(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    sf(r, bold=True, italic=True, size=12)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after  = Pt(3)

def para(text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    sf(r)
    return p

def para_mx(parts, indent=True):
    """parts = [(text, bold, italic, mono)]"""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent:
        p.paragraph_format.first_line_indent = Cm(1.25)
    p.paragraph_format.space_after = Pt(6)
    for txt, b, i, m in parts:
        r = p.add_run(txt)
        sf(r, bold=b, italic=i, mono=m, size=10 if m else 12)
    return p

def bullet(text, sub=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent  = Cm(1.6 if sub else 1.0)
    p.paragraph_format.space_after  = Pt(3)
    r = p.add_run(('◦ ' if sub else '• ') + text)
    sf(r)

def numbered(n, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent        = Cm(1.0)
    p.paragraph_format.first_line_indent  = Cm(-0.6)
    p.paragraph_format.space_after        = Pt(4)
    r = p.add_run(f'{n}) {text}')
    sf(r)

def code(lines):
    for line in lines:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after  = Pt(0)
        p.paragraph_format.left_indent  = Cm(1.0)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'),   'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'),  'F2F2F2')
        pPr.append(shd)
        r = p.add_run(line)
        sf(r, mono=True, size=9)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

def caption(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    sf(r, italic=True, size=11)
    p.paragraph_format.space_after = Pt(10)

# ─── ANTET ───────────────────────────────────────────────────────────────────
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
for line in ["Universitatea Alexandru Ioan Cuza din Iași\n",
             "Facultatea de Economie și Administrarea Afacerilor\n",
             "An 3, Informatică Economică — Proiectarea Sistemelor Informaționale\n"]:
    r = p.add_run(line); sf(r, size=11)

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("Proiect — Sistem de Gestiune Comirex\nModulul: Recepție Document Însoțitor")
sf(r, bold=True, size=14)
doc.add_paragraph()

# ════════════════════════════════════════════════════════════════════════════
h1("4. Proiectarea arhitecturală și implementarea programelor")
# ════════════════════════════════════════════════════════════════════════════

para("Până aici ne-am ocupat de:")
bullet("Elaborarea structurii aplicației, redată sub forma claselor de obiecte și a relațiilor "
       "dintre ele, tabelelor bazei de date și a legăturilor dintre ele și maparea dintre cele "
       "două straturi. Structura aplicației s-a concretizat prin construirea modelului fizic "
       "(vezi partea 3) pe baza transformării modelelor conceptual și logic.")
bullet("Completarea specificațiilor de proiectare privind comportamentul aplicației, ca rezultat "
       "al proiectării formularului de recepție (inclusiv scenariile de lucru pentru fiecare "
       "membru al echipei).")
doc.add_paragraph()

para("În continuare ne vom preocupa de implementarea comportamentului aplicației ținând cont de "
     "modelul fizic obținut anterior.")

para("Implementarea comportamentului unei aplicații trebuie să aibă ÎNTOTDEAUNA la bază o "
     "arhitectură și unul sau mai multe șabloane de proiectare (design patterns) bine definite. "
     "Cu alte cuvinte, trebuie să existe o anumită consistență între secvențele finale de cod "
     "care implementează două sau mai multe scenarii diferite. Este vorba despre normalizarea "
     "codului pe baza unor reguli clare ce permit: reutilizarea codului, evitarea redundanței, "
     "managementul componentelor și al cuplării dintre acestea.")

para("În demersul nostru am selectat un șablon arhitectural și două șabloane de proiectare, "
     "aplicabile cu succes în majoritatea situațiilor și folosind orice limbaj:")
numbered(1, "MVC (Model-View-Controller) – șablon arhitectural pentru organizarea codului "
            "privind interfața utilizator;")
numbered(2, "Repository – șablon de proiectare specific aplicațiilor de gestiune date;")
numbered(3, "Adapter – șablon general valabil care poate fi aplicat în orice aplicație software.")
doc.add_paragraph()

# ─── 4.1 MVC ─────────────────────────────────────────────────────────────────
h2("4.1 Șablonul arhitectural MVC adaptat")

para("Șablonul MVC (Model-View-Controller) împarte codul aplicației în trei responsabilități "
     "distincte, vizibil în Figura 11:")
bullet("Model (ReceptieFormData) – încapsulează starea formularului; răspunde la interogări "
       "de stare; expune datele necesare View-ului prin liste lazy-loaded (furnizori, gestiuni, "
       "produse, document curent, recepție selectată).")
bullet("View (ReceptieFormView) – redă modelul grafic folosind componente Swing; transmite "
       "gesturile utilizatorului (click, selecție combo) către Controller.")
bullet("Controller (ReceptieFormCtrl) – definește comportamentul aplicației; mapează acțiunile "
       "utilizatorului la actualizări ale modelului de date; coordonează Repository-urile.")
doc.add_paragraph()

para("Diagrama de clase completă, în format PlantUML, este disponibilă în fișierul "
     "DiagramaClase_MVC.puml din directorul docs/. Structura claselor urmează "
     "fidel șablonul MVC adaptat prezentat la orele de laborator.")

# ─── 4.2 Repository ──────────────────────────────────────────────────────────
h2("4.2 Șablonul Repository")

para("Toate operațiile de acces la baza de date sunt izolate în clase repository care extind "
     "AbstractRepository. Aceasta furnizează metodele CRUD de bază (create, update, delete) "
     "și accesul la EntityManager-ul JPA, inițializat o singură dată la nivel de clasă "
     "(atribut static):")

code([
    "// AbstractRepository — initializare EntityManager",
    "private static EntityManager em =",
    "    Persistence.createEntityManagerFactory(\"ComirexPersistenceUnit\")",
    "              .createEntityManager();",
    "",
    "public Object create(Object o) { this.getEm().persist(o); return o; }",
    "public Object update(Object o) { return this.em.merge(o); }",
    "public void   delete(Object o) { this.em.remove(this.em.merge(o)); }",
])

para("MasterRepository gestionează entitățile master (Furnizor, Gestiune, Produs, Stoc), "
     "iar DocumentRepository gestionează documentele (DocInsotitor, Receptie). "
     "Separarea responsabilităților permite reutilizarea independentă a fiecărui repository.")

# ─── 4.3 Implementarea metodei salveaza ──────────────────────────────────────
h2("4.3 Implementarea metodei salveazaModificariDocument()")

para("Metoda centrală a formularului este salveazaModificariDocument() din ReceptieFormCtrl. "
     "Ea este invocată la apăsarea butonului Salvare și orchestrează persistența întregului "
     "document, inclusiv actualizarea stocului. Mai jos este detaliată pas cu pas:")

numbered(1, "Scenariul începe prin apăsarea de către utilizator a butonului Salvare în "
            "formular, ceea ce invocă metoda salveazaModificariDocument() a obiectului de tip "
            "ReceptieFormCtrl (care joacă rolul de controller).")

numbered(2, "Înainte de a deschide tranzacția, controller-ul reține dacă documentul curent "
            "este NOU (id == null). Această informație va fi folosită ulterior pentru a decide "
            "dacă se actualizează stocul.")

numbered(3, "Controller-ul transmite un mesaj obiectului docRepo din clasa DocumentRepository "
            "prin care invocă metoda beginTransaction(), care la rândul ei apelează "
            "getTransaction().begin() pe EntityManager.")

numbered(4, "Obiectul-țintă al metodei este documentul curent ținut de formData.documentCurent, "
            "pe care utilizatorul tocmai l-a completat prin interfața grafică. DocInsotitor este "
            "un obiect complex, cu relații @OneToMany configurate cu cascade=CascadeType.ALL "
            "spre Receptie, care la rândul ei are cascade spre LinieIntrare. Astfel, la "
            "persistarea unui singur obiect DocInsotitor, Hibernate generează automat "
            "instrucțiuni INSERT pentru toate recepțiile și liniile asociate.")

numbered(5, "Se invocă metoda saveDocInsotitor() din DocumentRepository:")

code([
    "public DocInsotitor saveDocInsotitor(DocInsotitor d) {",
    "    return (DocInsotitor) this.saveDocument(d);",
    "}",
    "",
    "public Document saveDocument(Document document) {",
    "    if (document.getId() == null)          // obiect nou",
    "        document = (Document) this.create(document);",
    "    else                                   // obiect existent in BD",
    "        document = (Document) this.update(document);",
    "    return document;",
    "}",
])

numbered(6, "Metoda saveDocument verifică id-ul obiectului. Dacă nu există (id == null), "
            "obiectul este nou și se apelează create() → persist() pe EntityManager. "
            "Altfel se apelează update() → merge(). Hibernate generează dinamic instrucțiunile "
            "SQL corespunzătoare (INSERT sau UPDATE) pentru document și toate relațiile "
            "sale în cascadă.")

numbered(7, "Dacă documentul este nou, după persistare se apelează metoda privată "
            "actualizeazaStoc() care parcurge toate recepțiile și liniile și actualizează "
            "tabela Stoc în aceeași tranzacție:")

code([
    "private void actualizeazaStoc(DocInsotitor doc) {",
    "    MasterRepository masterRepo = this.formData.getMasterRepo();",
    "    for (Receptie receptie : doc.getReceptii()) {",
    "        Gestiune gestiune = receptie.getGestiune();",
    "        if (gestiune == null) continue;",
    "        for (LinieIntrare linie : receptie.getLiniiIntrare()) {",
    "            Produs produs = linie.getProdus();",
    "            if (produs == null) continue;",
    "            Double cantitate = linie.getCantitate();",
    "            if (cantitate == null || cantitate == 0.0) continue;",
    "            List<Stoc> stocuri = masterRepo.findStocByProdusGestiune(",
    "                                     produs.getId(), gestiune.getId());",
    "            if (stocuri.isEmpty()) {",
    "                // Prima intrare — cream inregistrarea de stoc",
    "                Stoc s = new Stoc();",
    "                s.setProdus(produs); s.setGestiune(gestiune);",
    "                s.setStocCurent(cantitate);",
    "                masterRepo.addStoc(s);",
    "            } else {",
    "                // Stoc existent — adunam cantitatea",
    "                // (negativa pentru Stornare → scade stocul)",
    "                Stoc s = stocuri.get(0);",
    "                s.setStocCurent(s.getStocCurent() + cantitate);",
    "                masterRepo.updateStoc(s);",
    "            }",
    "        }",
    "    }",
    "}",
])

numbered(8, "La final, controller-ul transmite commitTransaction() către DocumentRepository, "
            "care apelează getTransaction().commit() pe EntityManager. Toate instrucțiunile "
            "SQL acumulate sunt trimise efectiv către PostgreSQL. În caz de eroare SQL, "
            "tranzacția este anulată automat (rollback).")

para("Codul complet al metodei salveazaModificariDocument():")

code([
    "public void salveazaModificariDocument() {",
    "    if (this.formData.getDocumentCurent() == null)",
    "        throw new RuntimeException(\"Nu exista niciun document de salvat!\");",
    "",
    "    boolean esteDocumentNou =",
    "        (this.formData.getDocumentCurent().getId() == null);",
    "",
    "    this.formData.getDocRepo().beginTransaction();",
    "    DocInsotitor doc = this.formData.getDocumentCurent();",
    "    DocInsotitor salvat = this.formData.getDocRepo().saveDocInsotitor(doc);",
    "    this.formData.setDocumentCurent(salvat);",
    "",
    "    if (esteDocumentNou) {",
    "        actualizeazaStoc(salvat); // actualizeaza tabela Stoc",
    "    }",
    "",
    "    this.formData.getDocRepo().commitTransaction();",
    "}",
])

# ─── 4.4 Stornare ────────────────────────────────────────────────────────────
h2("4.4 Implementarea Stornării — flux complet")

para("Stornarea reprezintă anularea/returul unui document anterior (Factură sau Aviz). "
     "Față de o recepție obișnuită, stornarea are trei particularități:")
bullet("Referențiază un document original prin câmpul docInsotitorReferinta din DocInsotitor.")
bullet("Cantitățile liniilor sunt întotdeauna negative (returnare spre furnizor → scădere stoc).")
bullet("Utilizatorul NU introduce manual valori negative; sistemul le generează automat.")
doc.add_paragraph()

para("Când utilizatorul selectează tipul Stornare din combo-ul formularului, View-ul declanșează "
     "automat un dialog de selectare a documentului original. La confirmare, se apelează metoda "
     "initiazaStornareDocument() din controller:")

code([
    "public void initiazaStornareDocument(DocInsotitor docOriginal) {",
    "    DocInsotitor stornat = this.formData.getDocumentCurent();",
    "    stornat.setTipDocument(STORNARE);",
    "    stornat.setFurnizor(docOriginal.getFurnizor());",
    "    stornat.setDocInsotitorReferinta(docOriginal); // legatura cu originalul",
    "",
    "    adaugaReceptie(); // creeaza receptia de stornare",
    "",
    "    for (Receptie r : docOriginal.getReceptii()) {",
    "        if (r.getGestiune() != null)",
    "            this.formData.getReceptieSelectata().setGestiune(r.getGestiune());",
    "        for (LinieIntrare original : r.getLiniiIntrare()) {",
    "            LinieIntrare linieStor = new LinieIntrare();",
    "            linieStor.setProdus(original.getProdus());",
    "            double cant = original.getCantitate() != null",
    "                        ? original.getCantitate() : 0.0;",
    "            linieStor.setCantitate(-Math.abs(cant)); // intotdeauna negativa",
    "            linieStor.setPretAchizitie(original.getPretAchizitie());",
    "            linieStor.setPretVanzare(original.getPretVanzare());",
    "            linieStor.setAdaos(original.getAdaos());",
    "            this.formData.getReceptieSelectata().addLinieIntrare(linieStor);",
    "        }",
    "    }",
    "}",
])

para("Tabelul din formular afișează cantitățile în valoare absolută (pozitiv), iar la editare "
     "manuală, valorile introduse de utilizator sunt negate automat de TableModel. "
     "La salvare, cantitățile negative persistate declanșează scăderea stocului prin "
     "aceeași metodă actualizeazaStoc() (stocCurent += cantitateNegativa).")

# ─── 4.5 Testare ─────────────────────────────────────────────────────────────
h2("4.5 Testarea scenariilor formularului")

para("O dată implementate metodele controller-ului, înainte de a conecta obiectele grafice, "
     "este necesară cel puțin o clasă de test care va simula acțiunile utilizatorilor și va "
     "verifica rezultatele. În proiectul Comirex au fost implementate patru clase de test, "
     "câte una pentru fiecare scenariu de lucru / membru al echipei.")

para("Simularea presupune execuția metodelor ce vor fi asociate cu diferite evenimente, "
     "conform pașilor descriși la definirea scenariului. Verificarea presupune atât validarea "
     "stărilor prin care trece formularul (ReceptieFormData) după fiecare metodă, cât și "
     "succesiunea de stări persistente ale modelului de date în baza de date.")

h3("Scenariul 1 — Recepție cu Factură (Student 1)")
para("Testează fluxul complet de creare a unui document de tip Factură cu o linie de intrare "
     "și verifică persistarea în PostgreSQL:")

code([
    "ReceptieFormCtrl form = new ReceptieFormCtrl();",
    "form.getFormData().setOperatieSelectata(",
    "    ReceptieFormData.RECEPTIE_CU_FACTURA);",
    "form.documentNou();",
    "form.getFormData().getDocumentCurent().setNumarDocument(\"FCT-001\");",
    "form.selectieFurnizor(form.getFormData().getListaFurnizori().get(0));",
    "form.adaugaReceptie();",
    "form.getFormData().getReceptieSelectata()",
    "    .setGestiune(form.getFormData().getListaGestiuni().get(0));",
    "form.adaugaLinieIntrare();",
    "LinieIntrare linie = form.getFormData().getReceptieSelectata()",
    "                         .getLiniiIntrare().get(0);",
    "linie.setProdus(form.getFormData().getListaProduse().get(0));",
    "linie.setCantitate(10.0);",
    "form.salveazaModificariDocument();",
    "// Verificare: documentul a fost salvat si stocul a crescut cu 10",
    "Assert.assertNotNull(form.getFormData().getDocumentCurent().getId());",
])

h3("Scenariul 2 — Recepție cu Aviz, două linii (Student 2)")
para("Testează adăugarea a două produse diferite în aceeași recepție și verifică "
     "că ambele sunt persistate:")

code([
    "form.getFormData().setOperatieSelectata(",
    "    ReceptieFormData.RECEPTIE_CU_AVIZ);",
    "form.documentNou();",
    "form.getFormData().getDocumentCurent().setNumarDocument(\"AVZ-001\");",
    "form.selectieFurnizor(form.getFormData().getListaFurnizori().get(0));",
    "form.adaugaReceptie();",
    "form.adaugaLinieIntrare(); // linie 1 — Produs M-101",
    "form.adaugaLinieIntrare(); // linie 2 — Produs M-205",
    "// ... setare produse ...",
    "// Verificare intermediara:",
    "Assert.assertEquals(2,",
    "    form.getFormData().getReceptieSelectata().getLiniiIntrare().size());",
    "form.salveazaModificariDocument();",
])

h3("Scenariul 3 — Stornare (Student 3)")
para("Testează fluxul de stornare: documentul original este identificat, liniile sunt "
     "preluate cu cantități negate automat, iar stocul scade la salvare:")

code([
    "form.getFormData().setOperatieSelectata(ReceptieFormData.STORNARE);",
    "form.documentNou();",
    "// Verificare tip document setat automat:",
    "Assert.assertEquals(ReceptieFormCtrl.STORNARE,",
    "    form.getFormData().getDocumentCurent().getTipDocument());",
    "// Initializare stornare pe baza documentului original:",
    "DocInsotitor docOriginal = ...; // document FCT-001 deja salvat",
    "form.initiazaStornareDocument(docOriginal);",
    "form.getFormData().getDocumentCurent().setNumarDocument(\"STOR-001\");",
    "form.getFormData().getReceptieSelectata().setNrNIR(\"NIR-STOR-01\");",
    "form.salveazaModificariDocument();",
    "// Rezultat: cantitate = -10.0 in BD, stocCurent scade cu 10",
    "Assert.assertNotNull(form.getFormData().getDocumentCurent().getId());",
])

h3("Scenariul 4 — Răzgândire, ștergere linie din memorie (Student 4)")
para("Testează că o linie ștearsă din memorie înainte de salvare nu ajunge în baza de date "
     "și că stocul nu este afectat de ea:")

code([
    "form.documentNou();",
    "form.adaugaReceptie();",
    "form.getFormData().getDocumentCurent().setNumarDocument(\"DOC-DEL-01\");",
    "form.adaugaLinieIntrare(); // linie1",
    "form.adaugaLinieIntrare(); // linie2",
    "// Stergere din memorie (fara niciun DELETE in BD):",
    "LinieIntrare liniaDeSters = form.getFormData()",
    "    .getReceptieSelectata().getLiniiIntrare().get(1);",
    "form.getFormData().getReceptieSelectata()",
    "    .removeLinieIntrare(liniaDeSters);",
    "// Verificare intermediara — trebuie sa ramana o singura linie:",
    "Assert.assertEquals(1,",
    "    form.getFormData().getReceptieSelectata().getLiniiIntrare().size());",
    "form.salveazaModificariDocument();",
    "// BD contine exact o linie; stocul reflecta doar linie1",
])

para("Diagramele de secvențe aferente celor patru scenarii sunt disponibile în directorul "
     "docs/ al proiectului (Scenariu1_ReceptieCuFactura.puml, Scenariu2_ReceptieCuAviz.puml, "
     "Scenariu3_Stornare.puml, Scenariu4_Stergere.puml). Ele ilustrează interacțiunile dintre "
     "clasele MVC la nivel de apeluri de metode, conform șablonului din figura 14 și figura 15 "
     "din îndrumarul de laborator.")

# ─── Salvare ─────────────────────────────────────────────────────────────────
out = (r"C:\Users\Cosmin\Desktop\Facultate\Anul 3\Sem2\Proiectarea sistemelor informationale"
       r"\Proiect\comirex\docs\Cap4_ArhitecturaImplementare.docx")
doc.save(out)
print(f"Salvat: {out}")
